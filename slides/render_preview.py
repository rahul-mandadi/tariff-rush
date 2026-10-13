"""Render the saved .pptx to PNGs so the deck can be eyeballed without PowerPoint.

This reads the actual file rather than re-drawing from the build script, so a layout mistake
in the deck shows up here instead of being hidden by a parallel preview. It is a preview only:
PowerPoint's text layout will differ slightly, and this is not the submission artifact.

Usage:  python3 slides/render_preview.py
"""
import os
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

EMU = 914400.0
SCALE = 110          # pixels per inch
OUT = "slides/preview"
# the deck is set in Avenir Next, so the preview has to be too: Arial is wider and made
# every slide look more crowded here than it does in PowerPoint
AVENIR = "/System/Library/Fonts/Avenir Next.ttc"
IDX = {"reg": 7, "bold": 0, "ital": 4}       # faces inside the .ttc
F_MON = "/System/Library/Fonts/Supplemental/Courier New.ttf"
os.makedirs(OUT, exist_ok=True)


def font(size_pt, bold=False, italic=False, mono=False):
    px = max(8, int(size_pt * SCALE / 72))
    try:
        if mono:
            return ImageFont.truetype(F_MON, px)
        i = IDX["bold"] if bold else (IDX["ital"] if italic else IDX["reg"])
        return ImageFont.truetype(AVENIR, px, index=i)
    except OSError:
        return ImageFont.load_default()


def wrap(draw, text, fnt, maxw):
    out = []
    for para in text.split("\n"):
        if not para:
            out.append("")
            continue
        words, line = para.split(" "), ""
        for w in words:
            t = (line + " " + w).strip()
            if draw.textlength(t, font=fnt) <= maxw or not line:
                line = t
            else:
                out.append(line)
                line = w
        out.append(line)
    return out


prs = Presentation("slides/Mandadi_Khalsa_Presentation1.pptx")
W = int(prs.slide_width / EMU * SCALE)
H = int(prs.slide_height / EMU * SCALE)
paths = []
for i, slide in enumerate(prs.slides, 1):
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    for sh in slide.shapes:
        if sh.left is None:
            continue
        x, y = int(sh.left / EMU * SCALE), int(sh.top / EMU * SCALE)
        w = int((sh.width or 0) / EMU * SCALE)
        h = int((sh.height or 0) / EMU * SCALE)
        if sh.shape_type == 13:                      # picture
            try:
                pic = Image.open(sh.image.blob and __import__("io").BytesIO(sh.image.blob))
                img.paste(pic.convert("RGB").resize((w, h), Image.LANCZOS), (x, y))
            except Exception as e:
                d.rectangle([x, y, x + w, y + h], outline="red")
                d.text((x + 6, y + 6), f"[image: {e}]", fill="red", font=font(10))
            continue
        if not sh.has_text_frame:
            continue
        cy = y
        for p in sh.text_frame.paragraphs:
            if not p.runs:
                cy += int(10 * SCALE / 72)
                continue
            r = p.runs[0]
            sz = r.font.size.pt if r.font.size else 18
            col = r.font.color.rgb if (r.font.color and r.font.color.type is not None) else None
            fill = "#%s" % str(col) if col else "#000000"
            fnt = font(sz, bold=bool(r.font.bold), italic=bool(r.font.italic),
                       mono=(r.font.name == "Courier New"))
            txt = "".join(run.text for run in p.runs)
            mono = (r.font.name == "Courier New")
            # monospace blocks are pre-aligned, so re-wrapping would destroy the columns
            lines = [txt] if mono else wrap(d, txt, fnt, w)
            lh = int(sz * SCALE / 72 * 1.28)
            for ln in lines:
                tw = d.textlength(ln, font=fnt)
                tx = x + (w - tw) / 2 if str(p.alignment) == "CENTER (2)" else x
                d.text((tx, cy), ln, fill=fill, font=fnt)
                cy += lh
    p = f"{OUT}/slide{i:02d}.png"
    img.save(p)
    paths.append(p)
    print(f"  {p}")

# one contact sheet so the whole deck can be seen at once
cols, rows = 3, 4
tw, th = W // 3, H // 3
sheet = Image.new("RGB", (cols * tw + 40, rows * th + 40), "#eceff1")
for n, p in enumerate(paths):
    im = Image.open(p).resize((tw - 12, th - 12), Image.LANCZOS)
    sheet.paste(im, (20 + (n % cols) * tw, 20 + (n // cols) * th))
sheet.save(f"{OUT}/contact_sheet.png")
print(f"  {OUT}/contact_sheet.png")
